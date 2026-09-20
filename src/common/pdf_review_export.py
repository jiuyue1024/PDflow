"""Export confirmed ReviewSession overlays as a new XLSX copy."""

from __future__ import annotations

import os
from copy import deepcopy
from typing import Any, Dict, Iterable, List, Tuple

from openpyxl import load_workbook


class ReviewedExportError(ValueError):
    """Raised when a reviewed copy cannot be safely produced."""


def _same_path(left: str, right: str) -> bool:
    return os.path.normcase(os.path.abspath(left)) == os.path.normcase(os.path.abspath(right))


def _sheet_name_for_page(page_number: Any) -> str:
    return f"第{page_number}页"


def _find_block(recovery_result: Dict[str, Any], block_id: str) -> Dict[str, Any]:
    matches = [
        block
        for block in recovery_result.get("blocks") or []
        if isinstance(block, dict) and str(block.get("block_id")) == str(block_id)
    ]
    if not matches:
        raise ReviewedExportError(f"Reviewed block was not found: {block_id}")
    if len(matches) > 1:
        raise ReviewedExportError(f"Reviewed block identity is ambiguous: {block_id}")
    return matches[0]


def _find_sheet(workbook, block: Dict[str, Any]) -> Any:
    page_number = block.get("page_number")
    if page_number not in (None, 0):
        name = _sheet_name_for_page(page_number)
        if name in workbook.sheetnames:
            return workbook[name]
        raise ReviewedExportError(f"Original workbook sheet was not found for page {page_number}.")

    page_sheets = {
        _sheet_name_for_page(number)
        for number in range(1, len(workbook.sheetnames) + 1)
    }
    candidates = [sheet for sheet in workbook.worksheets if sheet.title not in page_sheets]
    if len(candidates) == 1:
        return candidates[0]
    if len(workbook.worksheets) == 1:
        return workbook.worksheets[0]
    raise ReviewedExportError("The fallback recovery sheet could not be identified safely.")


def _normalise_coordinate(coordinate: Any) -> Tuple[int, int]:
    if isinstance(coordinate, (tuple, list)) and len(coordinate) == 2:
        return int(coordinate[0]), int(coordinate[1])
    raise ReviewedExportError(f"Invalid reviewed cell coordinate: {coordinate!r}")


def _overlay_one(workbook, recovery_result: Dict[str, Any], reviewed_data: Dict[str, Any]) -> None:
    if reviewed_data.get("reviewed") is not True:
        raise ReviewedExportError("Only confirmed reviewed data can be exported.")
    block_id = reviewed_data.get("block_id")
    if not block_id:
        raise ReviewedExportError("Reviewed data has no stable block identity.")
    block = _find_block(recovery_result, str(block_id))
    page_number = block.get("page_number")
    same_page_blocks = [
        candidate
        for candidate in recovery_result.get("blocks") or []
        if isinstance(candidate, dict) and candidate.get("page_number") == page_number
    ]
    if len(same_page_blocks) > 1:
        raise ReviewedExportError(
            f"Reviewed block region is ambiguous for page {page_number}."
        )
    rows = reviewed_data.get("rows")
    if not isinstance(rows, list):
        raise ReviewedExportError("Reviewed data has no row payload.")
    sheet = _find_sheet(workbook, block)
    original_rows = block.get("rows") or []
    for coordinate in reviewed_data.get("changed_cells") or []:
        row, column = _normalise_coordinate(coordinate)
        if row < 0 or row >= len(rows) or column < 0 or column >= len(rows[row]):
            raise ReviewedExportError("Reviewed cell is outside the recovered table.")
        if row >= len(original_rows) or column >= len(original_rows[row]):
            raise ReviewedExportError("Reviewed cell is outside the original recovery data.")
        sheet.cell(row=row + 1, column=column + 1).value = rows[row][column]


def export_reviewed_xlsx(
    original_output_path: str,
    destination_path: str,
    recovery_result: Dict[str, Any],
    reviewed_data: Iterable[Dict[str, Any]],
) -> str:
    """Write a corrected copy without changing the source workbook or sidecar."""
    if not isinstance(recovery_result, dict):
        raise ReviewedExportError("A recovery result is required for reviewed export.")
    if not original_output_path or not os.path.isfile(original_output_path):
        raise ReviewedExportError("The original XLSX output was not found.")
    if not destination_path or _same_path(original_output_path, destination_path):
        raise ReviewedExportError("Choose a different path from the original output.")
    if os.path.exists(destination_path):
        raise ReviewedExportError("The selected destination already exists.")
    if not str(destination_path).lower().endswith(".xlsx"):
        raise ReviewedExportError("Reviewed export supports XLSX files only.")

    overlays = [item for item in reviewed_data if isinstance(item, dict)]
    if not overlays:
        raise ReviewedExportError("No confirmed reviewed data is available.")
    workbook = load_workbook(original_output_path)
    try:
        for overlay in overlays:
            _overlay_one(workbook, recovery_result, overlay)
        parent = os.path.dirname(os.path.abspath(destination_path))
        if parent:
            os.makedirs(parent, exist_ok=True)
        workbook.save(destination_path)
    finally:
        workbook.close()
    return destination_path
