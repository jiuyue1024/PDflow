import copy
import os
import shutil
import unittest
import uuid
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from src.common.pdf_review_export import ReviewedExportError, export_reviewed_xlsx
from src.common.pdf_review_session import create_review_session


def _recovery_result():
    return {
        "reliability_band": "review_required",
        "blocks": [
            {
                "block_id": "table-1",
                "page_number": 1,
                "rows": [["Name", "Age"], ["Alice", 30]],
                "cells": [
                    {"row_index": 0, "column_index": 0, "source_page": 1},
                    {"row_index": 0, "column_index": 1, "source_page": 1},
                    {"row_index": 1, "column_index": 0, "source_page": 1},
                    {"row_index": 1, "column_index": 1, "source_page": 1},
                ],
            },
            {
                "block_id": "table-2",
                "page_number": 2,
                "rows": [["Product", "Qty"], ["Pen", 2]],
                "cells": [],
            },
        ],
        "issues": [],
    }


def _make_workbook(path):
    workbook = Workbook()
    first = workbook.active
    first.title = "第1页"
    for row in [["Name", "Age"], ["Alice", 30]]:
        first.append(row)
    second = workbook.create_sheet("第2页")
    for row in [["Product", "Qty"], ["Pen", 2]]:
        second.append(row)
    workbook.save(path)
    workbook.close()


class _TestDirectory:
    def __init__(self, path):
        self.name = path

    def cleanup(self):
        shutil.rmtree(self.name, ignore_errors=True)


class DesktopVNextSlice4BTests(unittest.TestCase):
    def _paths(self):
        path = os.path.join(os.getcwd(), f".slice4b_test_{uuid.uuid4().hex}")
        os.makedirs(path, exist_ok=False)
        temp = _TestDirectory(path)
        original = os.path.join(path, "result.xlsx")
        _make_workbook(original)
        return temp, original, os.path.join(path, "result_reviewed.xlsx")

    def _reviewed(self, recovery=None, block_index=0, edits=None):
        recovery = recovery or _recovery_result()
        session = create_review_session(recovery, block_index)
        for row, column, value in edits or []:
            session.set_value(row, column, value)
        return session, recovery

    def test_confirmed_single_cell_edit_creates_copy_and_preserves_original(self):
        temp, original, destination = self._paths()
        try:
            session, recovery = self._reviewed(edits=[(1, 1, "31")])
            reviewed = session.apply()
            with open(original, "rb") as handle:
                before = handle.read()
            export_reviewed_xlsx(original, destination, recovery, [reviewed])
            with open(original, "rb") as handle:
                self.assertEqual(before, handle.read())
            workbook = load_workbook(destination, data_only=False)
            self.assertEqual("31", workbook["第1页"]["B2"].value)
            workbook.close()
        finally:
            temp.cleanup()

    def test_multiple_confirmed_edits_and_unchanged_cells(self):
        temp, original, destination = self._paths()
        try:
            session, recovery = self._reviewed(edits=[(0, 0, "Full name"), (1, 1, "31")])
            export_reviewed_xlsx(original, destination, recovery, [session.apply()])
            workbook = load_workbook(destination)
            self.assertEqual("Full name", workbook["第1页"]["A1"].value)
            self.assertEqual("31", workbook["第1页"]["B2"].value)
            self.assertEqual("Alice", workbook["第1页"]["A2"].value)
            workbook.close()
        finally:
            temp.cleanup()

    def test_unconfirmed_working_edit_never_exports(self):
        temp, original, destination = self._paths()
        try:
            session, recovery = self._reviewed(edits=[(1, 1, "31")])
            with self.assertRaises(ReviewedExportError):
                export_reviewed_xlsx(original, destination, recovery, [])
            self.assertFalse(os.path.exists(destination))
        finally:
            temp.cleanup()

    def test_cancelled_edit_never_exports(self):
        temp, original, destination = self._paths()
        try:
            session, recovery = self._reviewed(edits=[(1, 1, "31")])
            session.cancel()
            with self.assertRaises(ReviewedExportError):
                export_reviewed_xlsx(original, destination, recovery, [session.reviewed_data] if session.reviewed_data else [])
        finally:
            temp.cleanup()

    def test_unreviewed_blocks_and_order_remain_unchanged(self):
        temp, original, destination = self._paths()
        try:
            session, recovery = self._reviewed(edits=[(1, 1, "31")])
            export_reviewed_xlsx(original, destination, recovery, [session.apply()])
            workbook = load_workbook(destination)
            self.assertEqual(["第1页", "第2页"], workbook.sheetnames)
            self.assertEqual(2, workbook["第2页"]["B2"].value)
            workbook.close()
        finally:
            temp.cleanup()

    def test_original_recovery_result_and_output_path_are_unchanged(self):
        temp, original, destination = self._paths()
        try:
            recovery = _recovery_result()
            before = copy.deepcopy(recovery)
            session = create_review_session(recovery)
            session.set_value(1, 1, "31")
            export_reviewed_xlsx(original, destination, recovery, [session.apply()])
            self.assertEqual(before, recovery)
            self.assertEqual(original, original)
            self.assertTrue(os.path.isfile(original))
        finally:
            temp.cleanup()

    def test_original_destination_and_existing_destination_are_protected(self):
        temp, original, destination = self._paths()
        try:
            session, recovery = self._reviewed(edits=[(1, 1, "31")])
            reviewed = session.apply()
            with self.assertRaises(ReviewedExportError):
                export_reviewed_xlsx(original, original, recovery, [reviewed])
            with open(destination, "wb") as handle:
                handle.write(b"existing")
            with self.assertRaises(ReviewedExportError):
                export_reviewed_xlsx(original, destination, recovery, [reviewed])
        finally:
            temp.cleanup()

    def test_no_edit_apply_is_a_safe_noop_copy(self):
        temp, original, destination = self._paths()
        try:
            session, recovery = self._reviewed()
            export_reviewed_xlsx(original, destination, recovery, [session.apply()])
            self.assertTrue(os.path.isfile(destination))
            source_book = load_workbook(original)
            copied_book = load_workbook(destination)
            self.assertEqual(source_book.sheetnames, copied_book.sheetnames)
            for sheet_name in source_book.sheetnames:
                source_values = [[cell.value for cell in row] for row in source_book[sheet_name].iter_rows()]
                copied_values = [[cell.value for cell in row] for row in copied_book[sheet_name].iter_rows()]
                self.assertEqual(source_values, copied_values)
            source_book.close()
            copied_book.close()
        finally:
            temp.cleanup()

    def test_unusable_and_legacy_results_have_no_export_payload(self):
        temp, original, destination = self._paths()
        try:
            with self.assertRaises(ReviewedExportError):
                export_reviewed_xlsx(original, destination, {"blocks": []}, [])
            self.assertIsNone(create_review_session(None))
        finally:
            temp.cleanup()

    def test_multiple_reviewed_blocks_overlay_by_stable_block_id(self):
        temp, original, destination = self._paths()
        try:
            recovery = _recovery_result()
            first, _ = self._reviewed(recovery, 0, [(1, 1, "31")])
            second, _ = self._reviewed(recovery, 1, [(1, 1, 3)])
            export_reviewed_xlsx(original, destination, recovery, [first.apply(), second.apply()])
            workbook = load_workbook(destination)
            self.assertEqual("31", workbook["第1页"]["B2"].value)
            self.assertEqual(3, workbook["第2页"]["B2"].value)
            workbook.close()
        finally:
            temp.cleanup()

    def test_duplicate_block_identity_fails_safely(self):
        temp, original, destination = self._paths()
        try:
            recovery = _recovery_result()
            recovery["blocks"].append(copy.deepcopy(recovery["blocks"][0]))
            session, _ = self._reviewed(recovery, 0, [(1, 1, "31")])
            with self.assertRaises(ReviewedExportError):
                export_reviewed_xlsx(original, destination, recovery, [session.apply()])
            self.assertFalse(os.path.exists(destination))
        finally:
            temp.cleanup()

    def test_distinct_blocks_on_one_page_fail_safely(self):
        temp, original, destination = self._paths()
        try:
            recovery = _recovery_result()
            recovery["blocks"][1]["page_number"] = 1
            session, _ = self._reviewed(recovery, 0, [(1, 1, "31")])
            with self.assertRaises(ReviewedExportError):
                export_reviewed_xlsx(original, destination, recovery, [session.apply()])
            self.assertFalse(os.path.exists(destination))
        finally:
            temp.cleanup()

    def test_ui_apply_and_save_reviewed_copy(self):
        try:
            from PySide6.QtCore import QTimer
            from PySide6.QtWidgets import QApplication
            from pages.convert_page import ConvertPage, RecoveryReviewDialog
        except ModuleNotFoundError as exc:
            self.skipTest(f"Desktop UI dependencies unavailable: {exc}")
        temp, original, destination = self._paths()
        try:
            app = QApplication.instance() or QApplication([])
            page = ConvertPage()
            page._selected_type = "pdf_excel"
            page._results = [{"status": "ok", "output": original, "recovery_result": _recovery_result()}]
            page._apply_recovery_result_view()

            def apply_review():
                dialog = QApplication.activeModalWidget()
                self.assertIsInstance(dialog, RecoveryReviewDialog)
                dialog._table.item(1, 1).setText("31")
                dialog._apply_button.click()

            QTimer.singleShot(50, apply_review)
            page._btn_result_review.click()
            self.assertFalse(page._btn_reviewed_export.isHidden())
            with patch("pages.convert_page.QFileDialog.getSaveFileName", return_value=(destination, "Excel files (*.xlsx)")):
                page._btn_reviewed_export.click()
            workbook = load_workbook(destination)
            self.assertEqual("31", workbook["第1页"]["B2"].value)
            workbook.close()
            self.assertIn("Reviewed copy saved.", page._lbl_result_session.text())
            page.deleteLater()
        finally:
            temp.cleanup()


if __name__ == "__main__":
    unittest.main()
